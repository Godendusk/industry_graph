"""Offline presentation replay using existing report history; no model calls."""
import json
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from report_generation.word_export_agent import export_report_docx


def replay(history_path, output_dir):
    history = json.loads(Path(history_path).read_text())
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    rows = []
    for stamp, name in [('162226', 'low_altitude'), ('165322', 'embodied')]:
        source = next(row for row in history if stamp in row.get('preparedDocxPath', ''))
        result = export_report_docx(
            source['title'], source['abstractText'], source['bodySections'],
            industry=source['industry'], output_dir=str(output_dir / name),
            references=source['references'], adjustment_notes=source['adjustmentNotes'],
            report_scope=source['reportScope'], review=source['review'],
            summary_review=source['summaryReview'])
        assert result['status'] == 'success', result
        target = Path(result['docx_path'])
        if not target.is_absolute():
            target = ROOT / target
        paragraphs = [p.text for p in Document(target).paragraphs if p.text]
        original_path = Path(source['preparedDocxPath'])
        if not original_path.is_absolute():
            original_path = Path(history_path).resolve().parents[2] / original_path
        old = [p.text for p in Document(original_path).paragraphs if p.text]
        old = old[:old.index('报告调整说明')]
        old = [p for p in old if not p.startswith('未完成稿：')]
        # Exact prefix checks title, abstract, every heading/body paragraph and conditions.
        assert paragraphs[:len(old)] == old, '正文或原有限定发生了改变'
        assert paragraphs[len(old)] == '资料与方法说明'
        assert not any(p == '报告调整说明' or p.startswith('未完成稿：') for p in paragraphs)
        assert result['delivery_status'] == source['deliveryStatus'] == 'incomplete'
        assert result['delivery_reasons'] and result['adjustment_notes']
        assert any('事实待确认' in n['reason'] for n in result['adjustment_notes'])
        rows.append({'report': name, 'docx_path': str(target),
                     'unchanged_content_paragraphs': len(old),
                     'delivery_status': result['delivery_status'],
                     'delivery_reasons': result['delivery_reasons'],
                     'internal_note_count': len(result['adjustment_notes']),
                     'method_notes': result['method_notes']})
    (output_dir / 'replay.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2))
    return rows


if __name__ == '__main__':
    print(json.dumps(replay(sys.argv[1], sys.argv[2]), ensure_ascii=False, indent=2))
