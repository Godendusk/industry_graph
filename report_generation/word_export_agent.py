"""Word export agent for finalized industry reports."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, List


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SUPPORTED_INDUSTRIES = {
    "ai": "人工智能",
}
DEFAULT_OUTPUT_DIR = "report_generation/outputs"


def export_report_docx(
    report_title: str,
    abstract_text: str,
    body_sections: list,
    industry: str = "ai",
    output_dir: str = DEFAULT_OUTPUT_DIR,
) -> dict:
    """Export the final report as a formatted DOCX file."""
    normalized_title = _safe_text(report_title)
    normalized_abstract = _safe_text(abstract_text)
    normalized_industry = _safe_text(industry) or "ai"

    if not normalized_title:
        return _error_response("report_title cannot be empty", normalized_industry, "")

    industry_name = SUPPORTED_INDUSTRIES.get(normalized_industry)
    if not industry_name:
        return _error_response(
            f"unsupported industry: {normalized_industry}",
            normalized_industry,
            "",
        )

    if not normalized_abstract:
        return _error_response("abstract_text cannot be empty", normalized_industry, industry_name)
    if not isinstance(body_sections, list):
        return _error_response("body_sections must be an array", normalized_industry, industry_name)
    if not body_sections:
        return _error_response(
            "body_sections must contain at least one item",
            normalized_industry,
            industry_name,
        )

    warnings: List[dict] = []
    cleaned_abstract, abstract_warnings = _clean_report_text(
        normalized_abstract,
        stage="word_export_abstract_cleanup",
    )
    warnings.extend(abstract_warnings)
    if not cleaned_abstract:
        return _error_response(
            "abstract_text is empty after cleanup",
            normalized_industry,
            industry_name,
        )

    normalized_sections = _normalize_body_sections(body_sections, warnings)
    if not normalized_sections:
        return _error_response(
            "body_sections must contain at least one non-empty body_text",
            normalized_industry,
            industry_name,
        )

    output_path = _resolve_output_path(output_dir, normalized_title)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        _write_docx(
            path=output_path,
            report_title=normalized_title,
            abstract_text=cleaned_abstract,
            body_sections=normalized_sections,
        )
    except ImportError as exc:
        return _error_response(
            f"python-docx is required for Word export; run in kunlun environment: {exc}",
            normalized_industry,
            industry_name,
        )

    return {
        "status": "success",
        "industry": normalized_industry,
        "industry_name": industry_name,
        "report_title": normalized_title,
        "docx_path": _display_path(output_path),
        "warnings": warnings,
    }


def _normalize_body_sections(body_sections: list, warnings: List[dict]) -> List[dict]:
    normalized: List[dict] = []
    fallback_level1_index = 0
    fallback_level2_index = 0
    known_level1_keys: set[str] = set()

    for index, section in enumerate(body_sections, 1):
        if not isinstance(section, dict):
            warnings.append(
                {
                    "stage": "word_export_body_cleanup",
                    "message": f"skipped invalid body section at index {index}",
                }
            )
            continue

        outline_id = _safe_text(section.get("outline_id"))
        body_text, cleanup_warnings = _clean_report_text(
            section.get("body_text"),
            stage="word_export_body_cleanup",
            outline_id=outline_id,
        )
        warnings.extend(cleanup_warnings)
        if not body_text:
            warnings.append(
                {
                    "outline_id": outline_id,
                    "stage": "word_export_body_cleanup",
                    "message": "skipped body section with empty body_text",
                }
            )
            continue

        parent_title = _safe_text(section.get("parent_level1_title"))
        parent_id = _safe_text(section.get("parent_level1_id"))
        parent_key = parent_id or parent_title
        if not parent_key:
            parent_key = f"fallback_level1_{index}"
        if parent_key not in known_level1_keys:
            known_level1_keys.add(parent_key)
            fallback_level1_index += 1
            fallback_level2_index = 0
        fallback_level2_index += 1

        normalized.append(
            {
                "parent_key": parent_key,
                "parent_level1_title": parent_title or f"正文一级标题{fallback_level1_index}",
                "title": _safe_text(section.get("title")) or f"二级标题{fallback_level2_index}",
                "paragraphs": _split_paragraphs(body_text),
            }
        )
    return normalized


def _write_docx(path: Path, report_title: str, abstract_text: str, body_sections: List[dict]) -> None:
    docx = _import_docx_modules()
    document = docx.Document()
    _configure_section(document, docx)
    _configure_styles(document, docx)
    _configure_footer(document, docx)

    title = document.add_paragraph(style="ReportTitle")
    _add_single_run(title, report_title, "方正小标宋简体", "Times New Roman", 22)

    abstract_heading = document.add_paragraph(style="ReportHeading1")
    _add_single_run(abstract_heading, "摘要", "黑体", "SimHei", 16)
    for paragraph_text in _split_paragraphs(abstract_text):
        paragraph = document.add_paragraph(style="ReportBody")
        _add_body_runs(paragraph, paragraph_text)

    current_parent_key = ""
    level1_index = 0
    level2_index = 0
    for section in body_sections:
        if section["parent_key"] != current_parent_key:
            current_parent_key = section["parent_key"]
            level1_index += 1
            level2_index = 0
            level1_title = f"{_chinese_number(level1_index)}、{section['parent_level1_title']}"
            paragraph = document.add_paragraph(style="ReportHeading1")
            _add_single_run(paragraph, level1_title, "黑体", "SimHei", 16)

        level2_index += 1
        level2_title = f"（{_chinese_number(level2_index)}）{section['title']}"
        paragraph = document.add_paragraph(style="ReportHeading2")
        _add_single_run(paragraph, level2_title, "楷体", "KaiTi", 16)

        for paragraph_text in section["paragraphs"]:
            paragraph = document.add_paragraph(style="ReportBody")
            _add_body_runs(paragraph, paragraph_text)

    document.save(path)


def _import_docx_modules():
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    class Modules:
        pass

    modules = Modules()
    modules.Document = Document
    modules.WD_SECTION = WD_SECTION
    modules.WD_STYLE_TYPE = WD_STYLE_TYPE
    modules.WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH
    modules.WD_LINE_SPACING = WD_LINE_SPACING
    modules.OxmlElement = OxmlElement
    modules.qn = qn
    modules.Mm = Mm
    modules.Pt = Pt
    return modules


def _configure_section(document: Any, docx: Any) -> None:
    section = document.sections[0]
    section.page_width = docx.Mm(210)
    section.page_height = docx.Mm(297)
    section.top_margin = docx.Mm(37)
    section.bottom_margin = docx.Mm(35)
    section.left_margin = docx.Mm(27)
    section.right_margin = docx.Mm(27)


def _configure_footer(document: Any, docx: Any) -> None:
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = docx.WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number_field(paragraph, docx)


def _configure_styles(document: Any, docx: Any) -> None:
    normal = document.styles["Normal"]
    _set_style_font(normal, "仿宋", "Times New Roman", 16, docx)
    normal.paragraph_format.line_spacing_rule = docx.WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = docx.Pt(28)

    title = _get_or_add_style(document, "ReportTitle", docx)
    _set_style_font(title, "方正小标宋简体", "Times New Roman", 22, docx)
    title.paragraph_format.alignment = docx.WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing_rule = docx.WD_LINE_SPACING.EXACTLY
    title.paragraph_format.line_spacing = docx.Pt(32)
    title.paragraph_format.space_before = docx.Pt(0)
    title.paragraph_format.space_after = docx.Pt(12)

    heading1 = _get_or_add_style(document, "ReportHeading1", docx)
    _set_style_font(heading1, "黑体", "SimHei", 16, docx)
    heading1.paragraph_format.line_spacing_rule = docx.WD_LINE_SPACING.EXACTLY
    heading1.paragraph_format.line_spacing = docx.Pt(28)
    heading1.paragraph_format.space_before = docx.Pt(12)
    heading1.paragraph_format.space_after = docx.Pt(6)

    heading2 = _get_or_add_style(document, "ReportHeading2", docx)
    _set_style_font(heading2, "楷体", "KaiTi", 16, docx)
    heading2.paragraph_format.line_spacing_rule = docx.WD_LINE_SPACING.EXACTLY
    heading2.paragraph_format.line_spacing = docx.Pt(28)
    heading2.paragraph_format.space_before = docx.Pt(6)
    heading2.paragraph_format.space_after = docx.Pt(0)

    body = _get_or_add_style(document, "ReportBody", docx)
    _set_style_font(body, "仿宋", "Times New Roman", 16, docx)
    body.paragraph_format.line_spacing_rule = docx.WD_LINE_SPACING.EXACTLY
    body.paragraph_format.line_spacing = docx.Pt(28)
    body.paragraph_format.space_before = docx.Pt(0)
    body.paragraph_format.space_after = docx.Pt(0)
    body.paragraph_format.first_line_indent = docx.Pt(32)
    _set_style_first_line_chars(body, docx, "200")


def _get_or_add_style(document: Any, name: str, docx: Any) -> Any:
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, docx.WD_STYLE_TYPE.PARAGRAPH)


def _set_style_font(style: Any, east_asia_font: str, latin_font: str, size_pt: int, docx: Any) -> None:
    style.font.name = latin_font
    style.font.size = docx.Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    _set_rfonts(rpr, docx, latin_font, latin_font, east_asia_font, latin_font)
    _set_lang(rpr, docx, "zh-CN", "zh-CN")


def _set_style_first_line_chars(style: Any, docx: Any, value: str) -> None:
    ppr = style.element.get_or_add_pPr()
    ind = ppr.find(docx.qn("w:ind"))
    if ind is None:
        ind = docx.OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(docx.qn("w:firstLine"), "640")
    ind.set(docx.qn("w:firstLineChars"), value)


def _add_single_run(paragraph: Any, text: str, east_asia_font: str, latin_font: str, size_pt: int) -> None:
    docx = _import_docx_modules()
    run = paragraph.add_run(text)
    _format_run(
        run,
        ascii_font=latin_font,
        hansi_font=latin_font,
        east_asia_font=east_asia_font,
        cs_font=latin_font,
        size_pt=size_pt,
        lang_val="zh-CN",
        east_asia_lang="zh-CN",
        hint="eastAsia",
        docx=docx,
    )


def _add_body_runs(paragraph: Any, text: str) -> None:
    docx = _import_docx_modules()
    for segment_kind, segment_text in _split_body_font_segments(text):
        run = paragraph.add_run(segment_text)
        if segment_kind == "latin":
            _format_run(
                run,
                ascii_font="Times New Roman",
                hansi_font="Times New Roman",
                east_asia_font="Times New Roman",
                cs_font="Times New Roman",
                size_pt=16,
                lang_val="en-US",
                east_asia_lang="en-US",
                bidi_lang="en-US",
                hint="default",
                docx=docx,
            )
        else:
            _format_run(
                run,
                ascii_font="Times New Roman",
                hansi_font="Times New Roman",
                east_asia_font="仿宋",
                cs_font="Times New Roman",
                size_pt=16,
                lang_val="zh-CN",
                east_asia_lang="zh-CN",
                hint="eastAsia",
                docx=docx,
            )


def _format_run(
    run: Any,
    ascii_font: str,
    hansi_font: str,
    east_asia_font: str,
    cs_font: str,
    size_pt: int,
    lang_val: str,
    east_asia_lang: str,
    docx: Any,
    bidi_lang: str = "",
    hint: str = "",
) -> None:
    run.font.name = ascii_font
    run.font.size = docx.Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    _set_rfonts(rpr, docx, ascii_font, hansi_font, east_asia_font, cs_font, hint=hint)
    _set_run_size(rpr, docx, size_pt)
    _set_lang(rpr, docx, lang_val, east_asia_lang, bidi_lang=bidi_lang)


def _set_run_size(rpr: Any, docx: Any, size_pt: int) -> None:
    half_points = str(int(size_pt * 2))
    for tag_name in ("w:sz", "w:szCs"):
        node = rpr.find(docx.qn(tag_name))
        if node is None:
            node = docx.OxmlElement(tag_name)
            rpr.append(node)
        node.set(docx.qn("w:val"), half_points)


def _set_rfonts(
    rpr: Any,
    docx: Any,
    ascii_font: str,
    hansi_font: str,
    east_asia_font: str,
    cs_font: str,
    hint: str = "",
) -> None:
    rfonts = rpr.find(docx.qn("w:rFonts"))
    if rfonts is None:
        rfonts = docx.OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(docx.qn("w:ascii"), ascii_font)
    rfonts.set(docx.qn("w:hAnsi"), hansi_font)
    rfonts.set(docx.qn("w:eastAsia"), east_asia_font)
    rfonts.set(docx.qn("w:cs"), cs_font)
    if hint:
        rfonts.set(docx.qn("w:hint"), hint)


def _set_lang(
    rpr: Any,
    docx: Any,
    lang_val: str,
    east_asia_lang: str,
    bidi_lang: str = "",
) -> None:
    lang = rpr.find(docx.qn("w:lang"))
    if lang is None:
        lang = docx.OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(docx.qn("w:val"), lang_val)
    lang.set(docx.qn("w:eastAsia"), east_asia_lang)
    if bidi_lang:
        lang.set(docx.qn("w:bidi"), bidi_lang)


def _add_page_number_field(paragraph: Any, docx: Any) -> None:
    run = paragraph.add_run()
    fld_begin = docx.OxmlElement("w:fldChar")
    fld_begin.set(docx.qn("w:fldCharType"), "begin")

    instr = docx.OxmlElement("w:instrText")
    instr.set(docx.qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_separate = docx.OxmlElement("w:fldChar")
    fld_separate.set(docx.qn("w:fldCharType"), "separate")

    fallback = docx.OxmlElement("w:t")
    fallback.text = "1"

    fld_end = docx.OxmlElement("w:fldChar")
    fld_end.set(docx.qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(fallback)
    run._r.append(fld_end)


def _split_body_font_segments(text: str) -> List[tuple[str, str]]:
    segments: List[tuple[str, str]] = []
    current_kind = ""
    current_chars: List[str] = []

    for char in _strip_invalid_xml_chars(text):
        kind = "latin" if _is_ascii_printable(char) else "east_asia"
        if current_kind and kind != current_kind:
            segments.append((current_kind, "".join(current_chars)))
            current_chars = []
        current_kind = kind
        current_chars.append(char)

    if current_chars:
        segments.append((current_kind, "".join(current_chars)))
    return segments


def _is_ascii_printable(char: str) -> bool:
    return "\x20" <= char <= "\x7e"


def _clean_report_text(value: Any, stage: str, outline_id: str = "") -> tuple[str, List[dict]]:
    warnings: List[dict] = []
    text = _safe_text(value)

    fenced_match = re.fullmatch(r"```(?:markdown|md|text)?\s*([\s\S]*?)\s*```", text, re.IGNORECASE)
    if fenced_match:
        text = fenced_match.group(1).strip()
        warnings.append(_warning(stage, "removed enclosing code fence", outline_id))

    cleaned = re.sub(r"\[(?:知识图谱|外部资料)\d+\]", "", text)
    if cleaned != text:
        text = cleaned
        warnings.append(_warning(stage, "removed explicit citation markers", outline_id))

    reference_heading = re.search(
        r"(?im)^\s*(?:#{1,6}\s*)?(?:【)?(?:参考文献|引用列表|资料来源|资料来源说明)(?:】)?[:：]?\s*$",
        text,
    )
    if reference_heading:
        text = text[: reference_heading.start()].rstrip()
        warnings.append(_warning(stage, "removed trailing reference/source section", outline_id))

    return _strip_invalid_xml_chars(text).strip(), warnings


def _split_paragraphs(text: str) -> List[str]:
    paragraphs = [
        item.strip()
        for item in re.split(r"(?:\r?\n\s*)+", _safe_text(text))
        if item.strip()
    ]
    return paragraphs or [_safe_text(text)]


def _resolve_output_path(output_dir: str, report_title: str) -> Path:
    base = Path(output_dir or DEFAULT_OUTPUT_DIR)
    if not base.is_absolute():
        base = PROJECT_ROOT / base
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{_sanitize_filename(report_title)}_{timestamp}.docx"
    return base / filename


def _sanitize_filename(value: str) -> str:
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", _safe_text(value))
    text = re.sub(r"\s+", "", text).strip(" ._")
    if not text:
        text = "industry_report"
    # Linux 文件名上限 255 字节，中文 UTF-8 占 3 字节
    # 预留时间戳(_20260727_012305=16) + 扩展名(.docx=5) = 21 字节
    # 按字节截断，确保 filename 总长不超过 255 字节
    max_title_bytes = 230  # 255 - 21 - 少量余量
    encoded = text.encode("utf-8")
    if len(encoded) > max_title_bytes:
        encoded = encoded[:max_title_bytes]
        # 避免截断到 UTF-8 多字节字符中间，回退到最后一个完整字符
        text = encoded.decode("utf-8", errors="ignore")
    return text


def _display_path(path: Path) -> str:
    try:
        return path.resolve().relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        return path.as_posix()


def _chinese_number(value: int) -> str:
    digits = "零一二三四五六七八九"
    if value <= 0:
        return str(value)
    if value < 10:
        return digits[value]
    if value == 10:
        return "十"
    if value < 20:
        return "十" + digits[value % 10]
    if value < 100:
        tens, ones = divmod(value, 10)
        return digits[tens] + "十" + (digits[ones] if ones else "")
    return str(value)


def _strip_invalid_xml_chars(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))


def _warning(stage: str, message: str, outline_id: str = "") -> dict:
    warning = {
        "stage": stage,
        "message": message,
    }
    if outline_id:
        warning["outline_id"] = outline_id
    return warning


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _error_response(message: str, industry: str, industry_name: str) -> dict:
    return {
        "status": "error",
        "industry": industry,
        "industry_name": industry_name,
        "message": message,
        "docx_path": "",
        "warnings": [],
    }
