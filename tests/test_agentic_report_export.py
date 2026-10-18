from pathlib import Path
from docx import Document
from report_generation.word_export_agent import export_report_docx


def section():
    return {'outline_id':'S1.1','parent_level1_title':'竞争格局','title':'企业布局','status':'success',
      'body_text':'甲公司2025年计划扩产。','citation_ids':[],
      'evidence_pack':{'requirements':[{'requirement':'2025年市场份额','status':'insufficient','reason':'没有统一口径'}]},
      'claim_verification':{'status':'needs_confirmation','body_text':'甲公司2025年计划扩产。',
                            'claims':[{'status':'unknown','quote':'甲公司2025年计划扩产。','reason':'资料日期未确认'}]}}


def test_export_returns_pending_notes_even_without_frontend_notes(tmp_path):
    result=export_report_docx('报告','摘要',[section()],output_dir=str(tmp_path))
    assert result['status']=='success'
    text='\n'.join(p.text for p in Document(Path(result['docx_path'])).paragraphs)
    assert '报告调整说明' not in text and '没有统一口径' not in text and '资料日期未确认' not in text
    notes = '\n'.join(n['reason'] for n in result['adjustment_notes'])
    assert '没有统一口径' in notes and '资料日期未确认' in notes
    assert result['delivery_status'] == 'incomplete'


def test_export_treats_mismatched_pass_as_invalidated(tmp_path):
    s=section();s['claim_verification']={'status':'completed','body_text':'旧正文','claims':[]}
    result=export_report_docx('报告','摘要',[s],output_dir=str(tmp_path))
    text='\n'.join(p.text for p in Document(Path(result['docx_path'])).paragraphs)
    assert '原核对结果失效' not in text
    assert any('原核对结果失效' in n['reason'] for n in result['adjustment_notes'])
    assert result['delivery_status'] == 'incomplete'


def test_export_never_hides_omitted_claims(tmp_path):
    s=section();s['claim_verification']={'status':'needs_confirmation','body_text':s['body_text'],'claims':[], 'omitted_claim_count':3}
    result=export_report_docx('报告','摘要',[s],output_dir=str(tmp_path))
    text='\n'.join(p.text for p in Document(Path(result['docx_path'])).paragraphs)
    assert '未覆盖' not in text
    assert any('3项' in n['reason'] and '未覆盖' in n['reason'] for n in result['adjustment_notes'])
    assert result['delivery_status'] == 'incomplete'


def test_reexport_replaces_pending_snapshot_after_facts_are_checked(tmp_path):
    s=section()
    s.pop('evidence_pack')
    s['claim_verification']={'status':'completed','body_text':s['body_text'],'claims':[]}
    result=export_report_docx('报告','摘要',[s],output_dir=str(tmp_path),
        review={'status':'completed'},summary_review={'status':'completed'},
        adjustment_notes=[{'kind':'agentic_pending','reason':'旧的事实待确认'},
                          {'kind':'research_limit','reason':'旧的资料缺口'},
                          {'kind':'body_edit','reason':'此前压缩重复背景'}])
    notes='\n'.join(n['reason'] for n in result['adjustment_notes'])
    assert '旧的事实待确认' not in notes and '旧的资料缺口' not in notes
    assert '此前压缩重复背景' in notes
    assert result['delivery_status']=='completed'
