from pathlib import Path

from docx import Document

from report_generation.word_export_agent import export_report_docx


def test_word_uses_final_body_and_returns_internal_notes_outside_document(tmp_path):
    result = export_report_docx(
        report_title='测试报告', abstract_text='最终正文摘要',
        body_sections=[{'outline_id':'S1.1','parent_level1_title':'现状','title':'布局',
                        'body_text':'审查后的正文。','status':'success','citation_ids':[]}],
        output_dir=str(tmp_path), adjustment_notes=[
            {'kind':'body_edit','outline_ids':['S1.1'],'reason':'压缩重复背景'}])
    assert result['status'] == 'success'
    text = '\n'.join(p.text for p in Document(Path(result['docx_path'])).paragraphs)
    assert '审查后的正文' in text
    assert '报告调整说明' not in text
    assert '未完成稿' not in text
    assert '压缩重复背景' not in text
    assert '资料与方法说明' in text
    assert result['delivery_status'] == 'incomplete'
    assert result['delivery_reasons']
    assert any(n['reason'] == '压缩重复背景' for n in result['adjustment_notes'])


def test_method_notes_use_only_cited_sources_and_preserve_local_conditions(tmp_path):
    condition = '满载30件且单件0.5千克时，测算成本为每件1.15元，不能等同于实际盈利。[C1]'
    result = export_report_docx('报告', '摘要', [{
        'outline_id': 'S1.1', 'title': '运营', 'status': 'success',
        'body_text': condition, 'citation_ids': ['C1'],
    }], output_dir=str(tmp_path), references=[
        {'citation_id': 'C1', 'title': '运营数据', 'retrieved_at': '2026-10-18T16:00:00+08:00', 'publish_date': '2024-05-01'},
        {'citation_id': 'C2', 'title': '未使用资料', 'retrieved_at': '2026-10-19T16:00:00+08:00'},
    ], report_scope={'status': 'completed', 'boundary': {
        'subjects': ['国内低空物流'], 'period': '2024—2026年',
        'classification': '面向企业决策的报告'},
        'remaining_gaps': ['内部检索仍缺某份材料']},
        review={'status': 'completed'}, summary_review={'status': 'completed'})
    text = '\n'.join(p.text for p in Document(result['docx_path']).paragraphs)
    assert condition in text
    assert '国内低空物流' in text and '2024—2026年' in text
    assert '资料检索截止日期：2026-10-18' in text
    assert '2026-10-19' not in text and '内部检索仍缺' not in text
    assert '面向企业决策的报告' not in text  # 不是统计口径，不能据字段名擅自改写。
    assert text.index('资料与方法说明') < text.index('参考资料')
    assert result['delivery_status'] == 'completed'
    assert any('内部检索仍缺' in n['reason'] for n in result['adjustment_notes'])


def test_method_notes_do_not_invent_cutoff_when_retrieval_date_missing(tmp_path):
    result = export_report_docx('报告', '摘要', [{
        'outline_id': 'S1.1', 'title': '现状', 'status': 'success', 'body_text': '研究正文。',
    }], output_dir=str(tmp_path))
    text = '\n'.join(p.text for p in Document(result['docx_path']).paragraphs)
    assert '资料检索截止日期' not in text
    assert '资料与方法说明' in text
